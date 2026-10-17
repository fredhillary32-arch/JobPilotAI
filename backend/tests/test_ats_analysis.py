"""ATS requirement extraction, evidence matching, scoring and CV document analysis."""
from datetime import UTC, date, datetime, timedelta
from io import BytesIO
from uuid import UUID, uuid4

import pytest
from docx import Document
from fastapi.testclient import TestClient
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen.canvas import Canvas
from sqlalchemy import select
from sqlalchemy.orm import Session, sessionmaker

from app.models import (
    CandidateEvidence,
    CandidateProfile,
    CandidateSkill,
    CandidateWorkExperience,
    JobListing,
    User,
)
from app.services.job_requirements import extract_requirements

API = "/api/v1/candidate-profiles"

SUPPORT_JOB = """IT Support Specialist

About us:
We are a fintech company. Our teams chat in Slack and love Figma.

What you'll do:
- Resolve hardware and software tickets in Jira Service Management
- Administer Microsoft 365 accounts for new starters

Requirements:
- 3+ years of experience in IT support or help desk roles
- Troubleshooting Windows and macOS laptops
- Bachelor's degree in Computer Science or equivalent experience

Nice to have:
- CompTIA A+ certification
- Experience with Intune is a plus

Location: Remote. You must be based in Kenya or Uganda.
Benefits: health insurance and a laptop budget."""


def create_profile(client: TestClient, **fields) -> str:
    payload = {"name": "Amina Wanjiru", "professional_headline": "IT Support Specialist",
               "summary": "Supports staff with laptops, accounts and service desk tickets.", **fields}
    response = client.post(API, json=payload)
    assert response.status_code == 201, response.text
    return response.json()["id"]


def create_job(client: TestClient, profile_id: str, description: str = SUPPORT_JOB, title: str = "IT Support Specialist") -> str:
    response = client.post(f"{API}/{profile_id}/job-descriptions",
                           json={"job_title": title, "company": "Acme Pay", "description": description})
    assert response.status_code == 201, response.text
    return response.json()["id"]


def add_experience(factory, profile_id: str, *, achievements: list[str], start=date(2020, 1, 1),
                   end=None, title="IT Support Specialist", employer="Savanna Telecom", description=None) -> UUID:
    with factory() as session:
        work = CandidateWorkExperience(
            candidate_profile_id=UUID(profile_id), employer=employer, job_title=title, start_date=start,
            end_date=end, current_role=end is None, description=description, achievements=achievements)
        session.add(work)
        session.commit()
        return work.id


def add_skills(factory, profile_id: str, *names: str) -> None:
    with factory() as session:
        session.add_all(CandidateSkill(candidate_profile_id=UUID(profile_id), name=name) for name in names)
        session.commit()


def upload(client: TestClient, profile_id: str, filename: str, content: bytes, media_type: str) -> str:
    response = client.post(f"{API}/{profile_id}/documents", files={"file": (filename, content, media_type)})
    assert response.status_code == 201, response.text
    document_id = response.json()["id"]
    extracted = client.post(f"{API}/{profile_id}/documents/{document_id}/extraction")
    assert extracted.status_code == 200, extracted.text
    return document_id


def analysis(client: TestClient, profile_id: str, job_id: str, document_id: str | None = None) -> dict:
    query = f"?source_document_id={document_id}" if document_id else ""
    response = client.get(f"{API}/{profile_id}/job-descriptions/{job_id}/ats-analysis{query}")
    assert response.status_code == 200, response.text
    return response.json()


def by_text(items: list[dict], fragment: str) -> dict:
    return next(item for item in items if fragment.casefold() in item["source_text"].casefold())


# --- Requirement extraction -------------------------------------------------

def test_requirements_are_classified_with_exact_source_text_and_offsets() -> None:
    result = extract_requirements(SUPPORT_JOB, company="Acme Pay")
    requirements = result["requirements"]
    for item in requirements:
        assert SUPPORT_JOB[item["source_start"]:item["source_end"]] == item["source_text"]

    essential = {item["source_text"] for item in requirements if item["importance"] == "essential"}
    preferred = [item for item in requirements if item["importance"] == "preferred"]
    assert "Troubleshooting Windows and macOS laptops" in essential
    assert any("CompTIA A+" in item["source_text"] and item["basis"] == "explicit" for item in preferred)
    experience = next(item for item in requirements if item["category"] == "experience")
    assert experience["min_years"] == 3 and experience["importance"] == "essential"
    education = next(item for item in requirements if item["category"] == "education")
    assert education["degree_level"] == 3 and education["allows_equivalent"] is True
    assert education["terms"] == ["computer science"]

    duties = [item for item in requirements if item["category"] == "responsibility"]
    assert len(duties) == 2 and all(item["importance"] == "not_scored" for item in duties)
    # A tool named only in the duties is an inferred preference, not a stated requirement.
    inferred = next(item for item in requirements if item["terms"] == ["microsoft 365"] and item["category"] == "tool")
    assert inferred["basis"] == "inferred" and inferred["importance"] == "preferred"

    eligibility = [item for item in requirements if item["category"] == "eligibility"]
    assert any("based in Kenya" in item["source_text"] for item in eligibility)
    assert all(item["importance"] == "not_scored" for item in eligibility)


def test_descriptive_text_is_not_treated_as_requirements_or_keywords() -> None:
    result = extract_requirements(SUPPORT_JOB, company="Acme Pay")
    texts = " ".join(item["source_text"] for item in result["requirements"])
    assert "health insurance" not in texts
    assert "Slack" not in texts and "Figma" not in texts
    keywords = {keyword["term"] for keyword in result["keywords"]}
    assert "slack" not in keywords and "figma" not in keywords
    assert {"windows", "macos"} <= keywords


def test_mixed_required_and_optional_wording_is_marked_ambiguous() -> None:
    result = extract_requirements("Python is required, but Go is a plus for this role.")
    item = result["requirements"][0]
    assert item["importance"] == "essential" and item["basis"] == "ambiguous"


def test_curly_apostrophe_headings_are_recognised() -> None:
    text = "What you’ll do\n- Own a book of customers\n\nWhat you’ll bring\n- Experience selling SaaS products\n- Salesforce reporting"
    result = extract_requirements(text)
    essential = [item["source_text"] for item in result["requirements"] if item["importance"] == "essential"]
    duties = [item["source_text"] for item in result["requirements"] if item["category"] == "responsibility"]
    assert essential == ["Experience selling SaaS products", "Salesforce reporting"]
    assert duties == ["Own a book of customers"]


def test_duty_sentences_that_mention_required_are_not_requirements() -> None:
    text = ("Ensure safety and outlet compliance. Maintain required operating standards across stores. "
            "Stock operations need to be governed end to end. You must hold a valid driving licence.")
    result = extract_requirements(text)
    essential = [item["source_text"] for item in result["requirements"] if item["importance"] == "essential"]
    assert essential == ["You must hold a valid driving licence."]
    assert {item["category"] for item in result["requirements"] if item["importance"] == "not_scored"} == {"responsibility"}


def test_prose_without_requirement_wording_is_not_a_requirement() -> None:
    result = extract_requirements("Our office is near the park. We value kindness and curiosity.")
    assert [item for item in result["requirements"] if item["importance"] in {"essential", "preferred"}] == []


# --- Matching and scoring --------------------------------------------------

def test_listed_skill_alone_is_partial_and_demonstrated_work_is_matched(api_client: TestClient, database) -> None:
    factory = database[1]
    profile_id = create_profile(api_client)
    add_skills(factory, profile_id, "Windows", "macOS")
    job_id = create_job(api_client, profile_id)
    listed_only = by_text(analysis(api_client, profile_id, job_id)["requirements"]["essential"], "Troubleshooting Windows")
    assert listed_only["status"] == "partial"
    assert "listed as a skill" in listed_only["explanation"]

    add_experience(factory, profile_id, achievements=["Troubleshot Windows and macOS laptops for 300 staff across two offices."])
    shown = by_text(analysis(api_client, profile_id, job_id)["requirements"]["essential"], "Troubleshooting Windows")
    assert shown["status"] == "matched"
    assert any(ref["kind"] == "demonstrated" and ref["source_type"] == "work_experience" for ref in shown["evidence"])


def test_unreviewed_evidence_is_uncertain_and_rejected_evidence_is_ignored(api_client: TestClient, database) -> None:
    factory = database[1]
    profile_id = create_profile(api_client)
    document_id = upload(api_client, profile_id, "cv.txt", b"Set up Intune for laptops.\n", "text/plain")
    job_id = create_job(api_client, profile_id)
    with factory() as session:
        session.add(CandidateEvidence(
            candidate_profile_id=UUID(profile_id), source_document_id=UUID(document_id), category="achievement",
            description="Enrolled company laptops in Intune", source_excerpt="Set up Intune for laptops.",
            confirmation_status="unreviewed"))
        session.commit()
    intune = by_text(analysis(api_client, profile_id, job_id)["requirements"]["preferred"], "Intune")
    assert intune["status"] == "uncertain"

    with factory() as session:
        evidence = session.scalar(select(CandidateEvidence))
        evidence.confirmation_status = "rejected"
        session.commit()
    assert by_text(analysis(api_client, profile_id, job_id)["requirements"]["preferred"], "Intune")["status"] == "missing"

    with factory() as session:
        session.scalar(select(CandidateEvidence)).confirmation_status = "candidate_confirmed"
        session.commit()
    assert by_text(analysis(api_client, profile_id, job_id)["requirements"]["preferred"], "Intune")["status"] == "matched"


def test_uploaded_cv_text_counts_only_as_uncertain(api_client: TestClient, database) -> None:
    profile_id = create_profile(api_client)
    document_id = upload(api_client, profile_id, "cv.txt",
                         b"Work Experience\nTroubleshooting Windows and macOS laptops daily.\n", "text/plain")
    job_id = create_job(api_client, profile_id)
    report = analysis(api_client, profile_id, job_id, document_id)
    item = by_text(report["requirements"]["essential"], "Troubleshooting Windows")
    assert item["status"] == "uncertain"
    assert any(ref["source_type"] == "source_document" for ref in item["evidence"])
    assert report["potential_score"] > report["score"]


def test_score_is_deterministic_explained_and_weights_sum_to_100(api_client: TestClient, database) -> None:
    factory = database[1]
    profile_id = create_profile(api_client)
    add_experience(factory, profile_id, start=date(2019, 1, 1), achievements=[
        "Resolved hardware and software tickets in Jira Service Management for 200 staff.",
        "Troubleshot Windows and macOS laptops and replaced failing hardware.",
    ])
    job_id = create_job(api_client, profile_id)
    first, second = analysis(api_client, profile_id, job_id), analysis(api_client, profile_id, job_id)
    assert first["score"] == second["score"] and first["score"] is not None
    components = first["score_explanation"]["components"]
    assert round(sum(c["effective_weight"] for c in components)) == 100
    assert round(sum(c["points"] for c in components)) == first["score"]
    assert first["score_explanation"]["limitations"]
    assert 0 <= first["score"] <= first["potential_score"] <= 100


def test_more_demonstrated_evidence_never_lowers_the_score(api_client: TestClient, database) -> None:
    factory = database[1]
    profile_id = create_profile(api_client)
    job_id = create_job(api_client, profile_id)
    before = analysis(api_client, profile_id, job_id)["score"]
    add_experience(factory, profile_id, start=date(2018, 1, 1), achievements=["Troubleshot Windows and macOS laptops for remote staff."])
    middle = analysis(api_client, profile_id, job_id)["score"]
    add_skills(factory, profile_id, "Intune")
    after = analysis(api_client, profile_id, job_id)["score"]
    assert before <= middle <= after
    assert before < after


def test_empty_profile_gets_a_report_with_missing_items_and_no_errors(api_client: TestClient) -> None:
    profile_id = create_profile(api_client, professional_headline=None, summary=None)
    job_id = create_job(api_client, profile_id)
    report = analysis(api_client, profile_id, job_id)
    assert report["score"] == 0
    assert all(item["status"] in {"missing", "uncertain"} for item in report["requirements"]["essential"])
    assert report["experience_summary"] == {"total_years": 0.0, "roles": 0}
    assert any(rec["category"] == "gap" for rec in report["recommendations"])


def test_experience_shortfall_is_a_screening_concern(api_client: TestClient, database) -> None:
    factory = database[1]
    profile_id = create_profile(api_client)
    add_experience(factory, profile_id, start=date.today() - timedelta(days=400), achievements=["Resolved service desk tickets for staff."])
    job_id = create_job(api_client, profile_id)
    report = analysis(api_client, profile_id, job_id)
    experience = next(item for item in report["requirements"]["essential"] if item["category"] == "experience")
    assert experience["status"] == "missing"
    assert any("3+" in concern["text"] or "years" in concern["text"] for concern in report["screening_concerns"])


def test_employer_exclusion_overrides_a_perfect_skills_match(api_client: TestClient, database) -> None:
    factory = database[1]
    profile_id = create_profile(api_client)
    restriction = "Applicants must reside in the United Kingdom."
    description = f"Requirements:\n- Python\n- SQL\n\n{restriction}"
    add_experience(factory, profile_id, start=date(2015, 1, 1), achievements=["Built Python and SQL reporting jobs for finance teams."])
    now = datetime.now(UTC)
    with factory.begin() as session:
        listing = JobListing(
            source_kind="greenhouse", source_key="syntheticco", source_listing_id=str(uuid4()),
            source_url="https://boards-api.greenhouse.io/v1/boards/syntheticco/jobs?content=true",
            application_url="https://job-boards.greenhouse.io/syntheticco/jobs/1", employer_name="Synthetic Co",
            title="Data Analyst", description_text=description, location_text="Remote",
            content_hash="0" * 64, dedupe_key="1" * 64, source_status="active",
            first_seen_at=now, last_seen_at=now, last_verified_at=now)
        session.add(listing)
        session.flush()
        listing_id = listing.id
    recorded = api_client.post(f"/api/v1/job-listings/{listing_id}/eligibility-assessments", json={
        "outcome": "not_eligible", "explanation": "UK residents only.", "quotation": restriction,
        "evidence_field": "description", "statement_scope": "role", "role_applicability": "direct",
        "geographic_scope": "other_countries"})
    assert recorded.status_code == 201, recorded.text
    imported = api_client.post(f"{API}/{profile_id}/job-descriptions/from-listing/{listing_id}")
    assert imported.status_code == 201, imported.text
    report = analysis(api_client, profile_id, imported.json()["id"])
    assert report["score"] >= 75
    assert report["verdict"] == "excluded_by_employer"
    assert report["eligibility_overrides_score"] is True
    assert report["recommendations"][0]["category"] == "eligibility"


def test_confirmed_residence_conflict_is_a_potential_disqualifier(api_client: TestClient) -> None:
    profile_id = create_profile(api_client)
    confirmed = api_client.put(f"{API}/{profile_id}/eligibility-details", json={
        "confirmed": True, "country_of_residence": "KE", "relocation_willingness": "not_willing"})
    assert confirmed.status_code == 200, confirmed.text
    job_id = create_job(api_client, profile_id, "Requirements:\n- Python\n\nYou must reside in Germany.")
    report = analysis(api_client, profile_id, job_id)
    assert report["verdict"] == "potential_disqualifiers"
    assert report["eligibility"]["disqualifiers"]
    assert report["eligibility_overrides_score"] is True


# --- CV document analysis --------------------------------------------------

def _docx_with_table_and_header() -> bytes:
    document = Document()
    document.sections[0].header.paragraphs[0].text = "amina@example.com | +254 700 000 000"
    document.add_paragraph("Amina Wanjiru")
    table = document.add_table(rows=3, cols=2)
    rows = [("Experience", "IT Support Specialist, Savanna Telecom, 2019 to 2023"),
            ("Skills", "Windows, macOS, Jira Service Management, Microsoft 365"),
            ("Education", "Bachelor of Science in Computer Science, 2018")]
    for row, (left, right) in zip(table.rows, rows):
        row.cells[0].text, row.cells[1].text = left, right
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _two_column_pdf() -> bytes:
    output = BytesIO()
    canvas = Canvas(output, pagesize=letter)
    y = 720
    for index in range(8):
        canvas.drawString(50, y, f"Skill item {index} Windows support")
        canvas.drawString(330, y, f"Experience line {index} Savanna Telecom service desk")
        y -= 40
    canvas.save()
    return output.getvalue()


def test_docx_layout_tables_and_header_contact_are_reported(api_client: TestClient) -> None:
    profile_id = create_profile(api_client)
    document_id = upload(api_client, profile_id, "cv.docx", _docx_with_table_and_header(),
                         "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    job_id = create_job(api_client, profile_id)
    cv = analysis(api_client, profile_id, job_id, document_id)["cv_document"]
    codes = {risk["code"]: risk["severity"] for risk in cv["formatting_risks"]}
    assert codes["layout_table"] == "high"
    assert codes["contact_in_header"] == "high"
    assert cv["status"] == "analyzed" and "text" not in cv


def test_two_column_pdf_is_reported(api_client: TestClient) -> None:
    profile_id = create_profile(api_client)
    document_id = upload(api_client, profile_id, "cv.pdf", _two_column_pdf(), "application/pdf")
    job_id = create_job(api_client, profile_id)
    risks = analysis(api_client, profile_id, job_id, document_id)["cv_document"]["formatting_risks"]
    assert any(risk["code"] == "columns" and risk["severity"] == "high" for risk in risks)


def test_plain_cv_reports_missing_headings_unconfirmed_figures_and_date_conflicts(api_client: TestClient, database) -> None:
    factory = database[1]
    profile_id = create_profile(api_client)
    add_experience(factory, profile_id, start=date(2021, 3, 1), achievements=["Resolved service desk tickets for staff."])
    text = ("Amina Wanjiru\nSUPERPOWERS\nSavanna Telecom 2017 to 2019\n"
            "- Cut ticket backlog by 45% in three months\n- Resolved service desk tickets for staff\n")
    document_id = upload(api_client, profile_id, "cv.txt", text.encode(), "text/plain")
    job_id = create_job(api_client, profile_id)
    cv = analysis(api_client, profile_id, job_id, document_id)["cv_document"]
    codes = {risk["code"] for risk in cv["formatting_risks"]}
    assert {"missing_sections", "no_email", "unusual_heading"} <= codes
    assert any("45" in item["reason"] for item in cv["unconfirmed_statements"])
    assert cv["contradictions"] and "2021" in cv["contradictions"][0]["explanation"]


def test_malformed_or_unextracted_documents_are_reported_not_crashed(api_client: TestClient) -> None:
    profile_id = create_profile(api_client)
    job_id = create_job(api_client, profile_id)
    broken = api_client.post(f"{API}/{profile_id}/documents",
                             files={"file": ("cv.docx", b"PK\x03\x04not really a docx",
                                             "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    if broken.status_code == 201:
        document_id = broken.json()["id"]
        api_client.post(f"{API}/{profile_id}/documents/{document_id}/extraction")
        cv = analysis(api_client, profile_id, job_id, document_id)["cv_document"]
        assert cv["status"] == "unavailable" and cv["error"]
    else:
        assert broken.status_code in {400, 415, 422}
    plain = api_client.post(f"{API}/{profile_id}/documents", files={"file": ("cv.txt", b"Plain CV text", "text/plain")})
    cv = analysis(api_client, profile_id, job_id, plain.json()["id"])["cv_document"]
    assert cv["status"] == "unavailable" and "Extract" in cv["error"]


# --- Access control ---------------------------------------------------------

def test_analysis_requires_a_session_and_ownership(api_client: TestClient, unauthenticated_client: TestClient, database) -> None:
    factory = database[1]
    profile_id = create_profile(api_client)
    job_id = create_job(api_client, profile_id)
    with factory() as session:
        other = User(email="someone-else@example.com", password_hash="$argon2id$test-fixture")
        session.add(other)
        session.flush()
        foreign = CandidateProfile(name="Other Person", user_id=other.id)
        session.add(foreign)
        session.commit()
        foreign_id = foreign.id
    assert api_client.get(f"{API}/{foreign_id}/job-descriptions/{job_id}/ats-analysis").status_code == 404
    assert api_client.get(f"{API}/{profile_id}/job-descriptions/{uuid4()}/ats-analysis").status_code == 404
    assert api_client.get(f"{API}/{profile_id}/job-descriptions/{job_id}/ats-analysis?source_document_id={uuid4()}").status_code == 404

    api_client.cookies.clear()
    assert unauthenticated_client.get(f"{API}/{profile_id}/job-descriptions/{job_id}/ats-analysis").status_code == 401


@pytest.mark.parametrize("description", ["", "   "])
def test_blank_job_descriptions_are_still_refused(api_client: TestClient, description: str) -> None:
    profile_id = create_profile(api_client)
    response = api_client.post(f"{API}/{profile_id}/job-descriptions", json={"job_title": "Role", "description": description})
    assert response.status_code == 422
