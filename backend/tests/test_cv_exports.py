from datetime import date
from io import BytesIO
from uuid import UUID, uuid4

import pymupdf
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session, sessionmaker

from app.models import (
    CandidateEducation,
    CandidateProfile,
    CandidateSkill,
    CandidateWorkExperience,
    TailoredCvDraft,
)

ACHIEVEMENTS = [
    "Built a Python service to process customer requests more reliably.",
    "Automated AWS deployment checks to prevent configuration errors.",
    "Reduced release delays by testing deployment workflows before launch.",
    "Designed monitoring dashboards that helped teams resolve service incidents.",
    "Improved database query performance by analyzing slow request patterns.",
]
JOB_DESCRIPTION = (
    "Platform engineer role. Required skills: AWS, Terraform. "
    "Build reliable services and improve deployment workflows."
)
SECTION_ORDER = [
    "Professional Summary",
    "Tools and Platforms",
    "Core Skills",
    "Professional Experience",
    "Education",
]


def create_valid_draft(
    client: TestClient,
    factory: sessionmaker[Session],
    *,
    candidate_name: str = "Zoë O’Connor",
    summary: str = "Experienced engineer who improves reliable service delivery.",
    achievements: list[str] | None = None,
) -> tuple[str, str, str, str]:
    profile_response = client.post(
        "/api/v1/candidate-profiles",
        json={
            "name": candidate_name,
            "professional_headline": "Platform Engineer",
            "summary": summary,
        },
    )
    assert profile_response.status_code == 201
    profile_id = profile_response.json()["id"]

    with factory() as session:
        profile = session.get(CandidateProfile, UUID(profile_id))
        assert profile is not None
        session.add_all(
            [
                CandidateWorkExperience(
                    candidate_profile_id=profile.id,
                    employer="Northstar Systems",
                    job_title="Platform Engineer",
                    start_date=date(2022, 1, 1),
                    current_role=True,
                    achievements=achievements or ACHIEVEMENTS,
                ),
                CandidateEducation(
                    candidate_profile_id=profile.id,
                    institution="Example Institute",
                    qualification="Bachelor of Science",
                    field_of_study="Computer Science",
                    start_date=date(2017, 1, 1),
                    end_date=date(2021, 1, 1),
                ),
                CandidateSkill(candidate_profile_id=profile.id, name="AWS"),
                CandidateSkill(candidate_profile_id=profile.id, name="Python"),
            ]
        )
        session.commit()

    job_response = client.post(
        f"/api/v1/candidate-profiles/{profile_id}/job-descriptions",
        json={
            "job_title": "Platform Engineer",
            "description": JOB_DESCRIPTION,
        },
    )
    assert job_response.status_code == 201
    job_id = job_response.json()["id"]
    draft_response = client.post(
        f"/api/v1/candidate-profiles/{profile_id}/job-descriptions/"
        f"{job_id}/cv-drafts"
    )
    assert draft_response.status_code == 201, draft_response.text
    return profile_id, job_id, draft_response.json()["id"], draft_response.json()["generated_content"]


def export_path(profile_id: str, job_id: str, draft_id: str) -> str:
    return (
        f"/api/v1/candidate-profiles/{profile_id}/job-descriptions/"
        f"{job_id}/cv-drafts/{draft_id}/exports"
    )


def test_pdf_and_docx_exports_preserve_section_order_bullets_and_unicode(
    api_client: TestClient,
    database,
) -> None:
    profile_id, job_id, draft_id, original_draft_text = create_valid_draft(
        api_client, database[1]
    )
    source_bytes = "Original candidate CV content.".encode()
    uploaded = api_client.post(
        f"/api/v1/candidate-profiles/{profile_id}/documents",
        files={"file": ("original.txt", source_bytes, "text/plain")},
        data={"document_kind": "source_cv"},
    )
    assert uploaded.status_code == 201

    exports = []
    for export_format in ("pdf", "docx"):
        response = api_client.post(
            export_path(profile_id, job_id, draft_id),
            json={"export_format": export_format},
        )
        assert response.status_code == 201, response.text
        metadata = response.json()
        exports.append(metadata)
        assert metadata["export_format"] == export_format
        assert metadata["filename"].startswith("jobpilot-cv-")
        assert metadata["filename"].endswith(f".{export_format}")
        assert "/" not in metadata["filename"] and "\\" not in metadata["filename"]
        assert "storage_key" not in response.text
        assert "private-documents" not in response.text

        download_url = (
            export_path(profile_id, job_id, draft_id)
            + f"/{metadata['id']}/download"
        )
        download = api_client.get(download_url)
        assert download.status_code == 200
        assert download.headers["content-type"].startswith(metadata["media_type"])
        assert download.headers["content-disposition"].endswith(
            f'filename="{metadata["filename"]}"'
        )

        if export_format == "pdf":
            pdf = pymupdf.open(stream=download.content, filetype="pdf")
            text = "\n".join(page.get_text() for page in pdf)
            assert metadata["page_count"] == pdf.page_count
            assert pdf.page_count == 1
            assert "Zoë O’Connor" in text
            assert "Professional Summary" in text
            assert text.index("Professional Summary") < text.index("Tools and Platforms")
            assert text.index("Tools and Platforms") < text.index("Core Skills")
            assert text.index("Core Skills") < text.index("Professional Experience")
            assert text.index("Professional Experience") < text.index("Education")
            assert text.count("•") == 5
            assert "—" not in text and "&" not in text
        else:
            document = Document(BytesIO(download.content))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            body = "\n".join(paragraphs)
            assert metadata["page_count"] is None
            headings = [paragraph for paragraph in paragraphs if paragraph in SECTION_ORDER]
            assert headings == SECTION_ORDER
            assert "Zoë O’Connor" in body
            assert "Professional Summary" in body
            assert "—" not in body and "&" not in body
            bullet_paragraphs = [
                paragraph
                for paragraph in document.paragraphs
                if paragraph.style.name == "List Bullet"
            ]
            assert len(bullet_paragraphs) == 5
            assert all(
                len(paragraph.text.split()) >= 6 for paragraph in bullet_paragraphs
            )

    assert exports[0]["id"] != exports[1]["id"]
    assert exports[0]["document_id"] != exports[1]["document_id"]
    assert api_client.get(
        f"/api/v1/candidate-profiles/{profile_id}/cv-drafts/{draft_id}"
    ).json()["generated_content"] == original_draft_text
    assert api_client.get(
        f"/api/v1/candidate-profiles/{profile_id}/documents/"
        f"{uploaded.json()['id']}/download"
    ).content == source_bytes
    listed = api_client.get(export_path(profile_id, job_id, draft_id))
    assert {entry["id"] for entry in listed.json()} == {
        entry["id"] for entry in exports
    }


def test_failed_drafts_cannot_be_exported(
    api_client: TestClient,
    database,
) -> None:
    profile = api_client.post(
        "/api/v1/candidate-profiles",
        json={
            "name": "Candidate",
            "professional_headline": "Platform Engineer",
            "summary": "Experienced engineer who improves reliable service delivery.",
        },
    ).json()
    profile_id = profile["id"]
    with database[1]() as session:
        profile_model = session.get(CandidateProfile, UUID(profile_id))
        assert profile_model is not None
        session.add(
            CandidateWorkExperience(
                candidate_profile_id=profile_model.id,
                employer="Northstar Systems",
                job_title="Platform Engineer",
                start_date=date(2022, 1, 1),
                current_role=True,
                achievements=ACHIEVEMENTS[:2],
            )
        )
        session.commit()
    job_response = api_client.post(
        f"/api/v1/candidate-profiles/{profile_id}/job-descriptions",
        json={"job_title": "Engineer", "description": JOB_DESCRIPTION},
    )
    job_id = job_response.json()["id"]
    # Two supported bullets: the draft is saved as incomplete, never padded.
    incomplete = api_client.post(
        f"/api/v1/candidate-profiles/{profile_id}/job-descriptions/"
        f"{job_id}/cv-drafts"
    )
    assert incomplete.status_code == 201 and incomplete.json()["status"] == "needs_review"
    export = api_client.post(
        export_path(profile_id, job_id, incomplete.json()["id"]),
        json={"export_format": "pdf"},
    )
    assert export.status_code == 409
    assert "incomplete" in export.json()["detail"]

    # A draft that failed validation has no content and is never exported.
    with database[1]() as session:
        failed = TailoredCvDraft(
            candidate_profile_id=UUID(profile_id), job_description_id=UUID(job_id), status="failed",
            generated_content=None, structured_content=None, source_references=[],
            validation_results={"is_valid": False, "errors": ["synthetic integrity error"], "gaps": [], "warnings": []},
        )
        session.add(failed)
        session.commit()
        failed_id = str(failed.id)
    export = api_client.post(
        export_path(profile_id, job_id, failed_id),
        json={"export_format": "pdf"},
    )
    assert export.status_code == 409
    assert "Only successfully generated" in export.json()["detail"]


def test_exports_reject_missing_and_mismatched_profile_job_draft_and_export_ids(
    api_client: TestClient,
    database,
) -> None:
    profile_id, job_id, draft_id, _ = create_valid_draft(api_client, database[1])
    export = api_client.post(
        export_path(profile_id, job_id, draft_id),
        json={"export_format": "pdf"},
    ).json()
    other_profile = api_client.post(
        "/api/v1/candidate-profiles", json={"name": "Other"}
    ).json()["id"]

    wrong_profile = api_client.post(
        export_path(other_profile, job_id, draft_id),
        json={"export_format": "pdf"},
    )
    wrong_job = api_client.post(
        export_path(profile_id, str(uuid4()), draft_id),
        json={"export_format": "pdf"},
    )
    wrong_draft = api_client.post(
        export_path(profile_id, job_id, str(uuid4())),
        json={"export_format": "pdf"},
    )
    invalid_format = api_client.post(
        export_path(profile_id, job_id, draft_id),
        json={"export_format": "txt"},
    )
    wrong_export_download = api_client.get(
        export_path(other_profile, job_id, draft_id)
        + f"/{export['id']}/download"
    )
    unknown_export_download = api_client.get(
        export_path(profile_id, job_id, draft_id)
        + f"/{uuid4()}/download"
    )

    assert wrong_profile.status_code == 404
    assert wrong_job.status_code == 404
    assert wrong_draft.status_code == 404
    assert invalid_format.status_code == 422
    assert wrong_export_download.status_code == 404
    assert unknown_export_download.status_code == 404


def test_tampered_draft_validation_or_source_is_rejected(
    api_client: TestClient,
    database,
) -> None:
    profile_id, job_id, draft_id, _ = create_valid_draft(api_client, database[1])
    with database[1]() as session:
        from app.models import TailoredCvDraft

        draft = session.get(TailoredCvDraft, UUID(draft_id))
        assert draft is not None
        draft.validation_results = {"is_valid": False, "errors": ["invalid"]}
        session.commit()

    response = api_client.post(
        export_path(profile_id, job_id, draft_id),
        json={"export_format": "docx"},
    )
    assert response.status_code == 409
    assert "did not pass content validation" in response.json()["detail"]


def test_long_cv_paginates_readably_and_reports_pdf_pages(
    api_client: TestClient,
    database,
) -> None:
    long_summary = " ".join(
        "Developed reliable services and improved platform operations."
        for _ in range(55)
    )[:4990]
    profile_id, job_id, draft_id, _ = create_valid_draft(
        api_client,
        database[1],
        summary=long_summary,
    )
    exported = api_client.post(
        export_path(profile_id, job_id, draft_id),
        json={"export_format": "pdf"},
    )
    assert exported.status_code == 201
    assert exported.json()["page_count"] > 1
    downloaded = api_client.get(
        export_path(profile_id, job_id, draft_id)
        + f"/{exported.json()['id']}/download"
    )
    pdf = pymupdf.open(stream=downloaded.content, filetype="pdf")
    assert pdf.page_count == exported.json()["page_count"]
    assert all(len(page.get_text()) > 0 for page in pdf)


def test_export_files_are_private_and_missing_file_is_reported(
    api_client: TestClient,
    database,
    tmp_path,
) -> None:
    profile_id, job_id, draft_id, _ = create_valid_draft(api_client, database[1])
    created = api_client.post(
        export_path(profile_id, job_id, draft_id),
        json={"export_format": "docx"},
    )
    metadata = created.json()
    file_path = tmp_path / "private-documents" / metadata["filename"]
    assert file_path.exists() is False
    storage_dir = tmp_path / "private-documents"
    stored_files = list(storage_dir.iterdir())
    assert len(stored_files) == 1
    assert stored_files[0].parent == storage_dir
    assert stored_files[0].suffix == ".docx"

    stored_files[0].unlink()
    missing = api_client.get(
        export_path(profile_id, job_id, draft_id)
        + f"/{metadata['id']}/download"
    )
    assert missing.status_code == 404
    assert str(tmp_path) not in missing.text


def test_missing_optional_contact_fields_do_not_prevent_export(
    api_client: TestClient,
    database,
) -> None:
    profile_id, job_id, draft_id, _ = create_valid_draft(api_client, database[1])
    export = api_client.post(
        export_path(profile_id, job_id, draft_id),
        json={"export_format": "pdf"},
    )
    assert export.status_code == 201
