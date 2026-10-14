from hashlib import sha256
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from uuid import uuid4
from zipfile import ZIP_DEFLATED, ZipFile

import pymupdf
import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.services.documents import LocalDocumentStorage
from app.services.extraction import ExtractionFailure, extract_document_text

PDF_TYPE = "application/pdf"
DOCX_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)


def create_profile(client: TestClient, name: str = "Extraction Owner") -> str:
    response = client.post("/api/v1/candidate-profiles", json={"name": name})
    assert response.status_code == 201
    return response.json()["id"]


def make_pdf(text: str = "CV evidence: improved service reliability.") -> bytes:
    document = pymupdf.open()
    page = document.new_page()
    if text:
        page.insert_text((72, 72), text)
    content = document.tobytes()
    document.close()
    return content


def make_docx() -> bytes:
    document = Document()
    document.add_paragraph("Professional summary")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Platform Engineer"
    table.cell(0, 1).text = "Reduced deployment time"
    document.add_paragraph("Education")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_malformed_docx() -> bytes:
    buffer = BytesIO()
    with ZipFile(buffer, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"<document")
    return buffer.getvalue()


def upload_document(
    client: TestClient,
    profile_id: str,
    filename: str,
    content_type: str,
    content: bytes,
) -> dict:
    response = client.post(
        f"/api/v1/candidate-profiles/{profile_id}/documents",
        files={"file": (filename, content, content_type)},
        data={"document_kind": "source_cv"},
    )
    assert response.status_code == 201, response.text
    return response.json()


def extract_via_api(
    client: TestClient, profile_id: str, document_id: str
) -> dict:
    return client.post(
        f"/api/v1/candidate-profiles/{profile_id}/documents/{document_id}/extraction"
    ).json()


def test_extracts_pdf_docx_and_utf8_text_without_changing_original(
    api_client: TestClient,
) -> None:
    profile_id = create_profile(api_client)
    fixtures = [
        ("source.pdf", PDF_TYPE, make_pdf(), "CV evidence: improved service reliability."),
        (
            "source.docx",
            DOCX_TYPE,
            make_docx(),
            "Professional summary\n\nPlatform Engineer\tReduced deployment time\n\nEducation",
        ),
        ("source.txt", "text/plain", "Résumé — Python developer.\n".encode(), "Résumé — Python developer."),
    ]

    for filename, content_type, content, expected_text in fixtures:
        source = upload_document(
            api_client, profile_id, filename, content_type, content
        )
        extracted = extract_via_api(api_client, profile_id, source["id"])
        assert extracted["status"] == "succeeded"
        assert extracted["character_count"] == len(expected_text)
        assert "storage_key" not in extracted
        assert "unverified_source_text" == extracted["trust_status"]

        text = api_client.get(
            f"/api/v1/candidate-profiles/{profile_id}/documents/"
            f"{source['id']}/extraction/text"
        )
        assert text.status_code == 200
        assert text.json()["text"] == expected_text

        original = api_client.get(
            f"/api/v1/candidate-profiles/{profile_id}/documents/"
            f"{source['id']}/download"
        )
        assert original.content == content


def test_no_text_malformed_encrypted_and_bad_utf8_are_reported_safely(
    api_client: TestClient,
) -> None:
    profile_id = create_profile(api_client)

    empty_pdf = pymupdf.open()
    empty_pdf.new_page()
    empty_pdf_bytes = empty_pdf.tobytes()
    empty_pdf.close()
    encrypted_pdf = pymupdf.open()
    encrypted_pdf.new_page()
    encrypted_bytes = encrypted_pdf.tobytes(
        encryption=pymupdf.PDF_ENCRYPT_AES_256,
        user_pw="test-only",
        owner_pw="test-owner",
    )
    encrypted_pdf.close()

    fixtures = [
        ("blank.pdf", PDF_TYPE, empty_pdf_bytes, "no_text", None),
        ("broken.pdf", PDF_TYPE, b"%PDF-1.7\nnot a valid PDF", "failed", "invalid_document"),
        ("locked.pdf", PDF_TYPE, encrypted_bytes, "failed", "encrypted_document"),
        ("bad-utf8.txt", "text/plain", b"a" * 4096 + b"\xff", "failed", "decoding_error"),
        ("broken.docx", DOCX_TYPE, make_malformed_docx(), "failed", "invalid_document"),
    ]
    for filename, content_type, content, expected_status, expected_error in fixtures:
        source = upload_document(
            api_client, profile_id, filename, content_type, content
        )
        result = extract_via_api(api_client, profile_id, source["id"])
        assert result["status"] == expected_status
        assert result["error_category"] == expected_error
        assert "Traceback" not in str(result)
        assert "storage_key" not in str(result)
        assert api_client.get(
            f"/api/v1/candidate-profiles/{profile_id}/documents/"
            f"{source['id']}/extraction/text"
        ).json()["text"] is None


def test_missing_source_path_and_size_limits_are_safe(
    api_client: TestClient,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from app.services import extraction

    profile_id = create_profile(api_client)
    source = upload_document(
        api_client, profile_id, "source.txt", "text/plain", b"safe source"
    )
    storage_root = tmp_path / "private-documents"
    stored = next(storage_root.iterdir())
    stored.unlink()
    result = extract_via_api(api_client, profile_id, source["id"])
    assert result["status"] == "failed"
    assert result["error_category"] == "missing_file"
    assert str(tmp_path) not in str(result)

    storage = LocalDocumentStorage(tmp_path / "limited")
    storage.storage_root.mkdir()
    data = b"too large"
    (storage.storage_root / "source.txt").write_bytes(data)
    monkeypatch.setattr(extraction, "MAX_DOCUMENT_SIZE_BYTES", 4)
    document = SimpleNamespace(
        media_type="text/plain",
        storage_key="source.txt",
        size_bytes=len(data),
        sha256=sha256(data).hexdigest(),
    )
    with pytest.raises(ExtractionFailure) as error:
        extract_document_text(document, storage)
    assert error.value.category == "resource_limit"

    unsafe_document = SimpleNamespace(
        media_type="text/plain",
        storage_key="../outside.txt",
        size_bytes=len(data),
        sha256=sha256(data).hexdigest(),
    )
    with pytest.raises(ExtractionFailure) as unsafe_error:
        extract_document_text(unsafe_document, storage)
    assert unsafe_error.value.category == "missing_file"

    monkeypatch.setattr(extraction, "MAX_DOCUMENT_SIZE_BYTES", 10 * 1024 * 1024)
    pdf_bytes = make_pdf()
    page_limited_document = SimpleNamespace(
        media_type=PDF_TYPE,
        storage_key="page-limited.pdf",
        size_bytes=len(pdf_bytes),
        sha256=sha256(pdf_bytes).hexdigest(),
    )
    (storage.storage_root / "page-limited.pdf").write_bytes(pdf_bytes)
    monkeypatch.setattr(extraction, "MAX_PDF_PAGES", 0)
    with pytest.raises(ExtractionFailure) as page_limit_error:
        extract_document_text(page_limited_document, storage)
    assert page_limit_error.value.category == "resource_limit"

    unknown_bytes = b"unsupported content"
    (storage.storage_root / "unknown.bin").write_bytes(unknown_bytes)
    unknown_document = SimpleNamespace(
        media_type="application/octet-stream",
        storage_key="unknown.bin",
        size_bytes=len(unknown_bytes),
        sha256=sha256(unknown_bytes).hexdigest(),
    )
    with pytest.raises(ExtractionFailure) as unsupported_error:
        extract_document_text(unknown_document, storage)
    assert unsupported_error.value.category == "unsupported_content"


def test_retry_updates_existing_extraction_row_and_detects_changed_source(
    api_client: TestClient,
    tmp_path: Path,
) -> None:
    profile_id = create_profile(api_client)
    source = upload_document(
        api_client, profile_id, "source.txt", "text/plain", b"Original evidence."
    )
    first = extract_via_api(api_client, profile_id, source["id"])
    assert first["status"] == "succeeded"

    root = tmp_path / "private-documents"
    stored = next(root.iterdir())
    stored.write_bytes(b"tampered source")
    second = extract_via_api(api_client, profile_id, source["id"])

    assert second["id"] == first["id"]
    assert second["status"] == "failed"
    assert second["error_category"] == "invalid_document"
    metadata = api_client.get(
        f"/api/v1/candidate-profiles/{profile_id}/documents/"
        f"{source['id']}/extraction"
    )
    assert metadata.json()["id"] == first["id"]


def test_evidence_requires_owned_document_and_matching_extracted_source(
    api_client: TestClient,
) -> None:
    owner_id = create_profile(api_client, "Evidence Owner")
    other_id = create_profile(api_client, "Other Candidate")
    source = upload_document(
        api_client, owner_id, "source.txt", "text/plain",
        b"Built a Python service that handled 100 requests per second.",
    )
    extract_via_api(api_client, owner_id, source["id"])
    payload = {
        "source_document_id": source["id"],
        "category": "achievement",
        "description": "Built a Python service.",
        "source_excerpt": "Built a Python service",
    }

    cross_profile = api_client.post(
        f"/api/v1/candidate-profiles/{other_id}/evidence", json=payload
    )
    invalid_excerpt = api_client.post(
        f"/api/v1/candidate-profiles/{owner_id}/evidence",
        json={**payload, "source_excerpt": "Not present in the source"},
    )
    created = api_client.post(
        f"/api/v1/candidate-profiles/{owner_id}/evidence", json=payload
    )

    assert cross_profile.status_code == 404
    assert invalid_excerpt.status_code == 422
    assert created.status_code == 201
    record = created.json()
    assert record["confirmation_status"] == "unreviewed"
    assert record["source_excerpt"] == "Built a Python service"
    assert record["source_start"] == 0
    assert record["source_end"] == len(record["source_excerpt"])

    evidence_path = (
        f"/api/v1/candidate-profiles/{owner_id}/evidence/{record['id']}"
    )
    null_status = api_client.patch(
        evidence_path, json={"confirmation_status": None}
    )
    confirmed = api_client.patch(
        evidence_path, json={"confirmation_status": "candidate_confirmed"}
    )
    rejected = api_client.patch(
        evidence_path, json={"confirmation_status": "rejected"}
    )
    wrong_owner = api_client.get(
        f"/api/v1/candidate-profiles/{other_id}/evidence/{record['id']}"
    )
    removed = api_client.delete(evidence_path)

    assert confirmed.json()["confirmation_status"] == "candidate_confirmed"
    assert null_status.status_code == 422
    assert rejected.json()["confirmation_status"] == "rejected"
    assert wrong_owner.status_code == 404
    assert removed.status_code == 204
    assert api_client.get(evidence_path).status_code == 404


def test_evidence_rejects_unknown_source_and_requires_extraction(
    api_client: TestClient,
) -> None:
    profile_id = create_profile(api_client)
    payload = {
        "source_document_id": str(uuid4()),
        "category": "technical_skill",
        "description": "Uses Python",
        "source_excerpt": "Python",
    }
    unknown = api_client.post(
        f"/api/v1/candidate-profiles/{profile_id}/evidence", json=payload
    )
    assert unknown.status_code == 404

    source = upload_document(
        api_client, profile_id, "source.txt", "text/plain", b"Python"
    )
    without_extraction = api_client.post(
        f"/api/v1/candidate-profiles/{profile_id}/evidence",
        json={**payload, "source_document_id": source["id"]},
    )
    assert without_extraction.status_code == 409
