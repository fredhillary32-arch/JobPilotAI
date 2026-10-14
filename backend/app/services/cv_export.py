from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date
from io import BytesIO
from pathlib import Path
from typing import Literal
from uuid import UUID, uuid4
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen.canvas import Canvas
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
)
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import (
    CandidateDocument,
    CandidateEducation,
    CandidateEvidence,
    CandidateProfile,
    CandidateSkill,
    CandidateWorkExperience,
    TailoredCvDraft,
    TailoredCvExport,
)
from app.schemas.cv_generation import CVSourceReference, CVStructuredContent
from app.services.cv_generation import render_cv
from app.services.cv_validation import validate_cv_content
from app.services.documents import DocumentStorage, DocumentStorageError

ExportFormat = Literal["pdf", "docx"]
MEDIA_TYPES: dict[ExportFormat, str] = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}
SPACE_PATTERN = re.compile(r"\s+")


class InvalidDraftForExport(ValueError):
    pass


class ExportFileUnavailable(OSError):
    pass


@dataclass(frozen=True)
class RenderedCv:
    content: bytes
    page_count: int | None
    media_type: str


def _clean(value: str) -> str:
    return SPACE_PATTERN.sub(" ", value).strip()


def _work_dates(experience: CandidateWorkExperience) -> str:
    start = experience.start_date.strftime("%b %Y")
    end = (
        "Present"
        if experience.current_role
        else experience.end_date.strftime("%b %Y")
        if experience.end_date
        else ""
    )
    return f"{start} to {end}" if end else start


def _assert_reference_matches(
    reference: CVSourceReference,
    profile: CandidateProfile,
    session: Session,
) -> None:
    if reference.source_type == "profile_field":
        fields = {
            "name": profile.name,
            "professional_headline": profile.professional_headline,
            "summary": profile.summary,
            "location": profile.location,
            "contact_email": profile.contact_email,
            "contact_phone": profile.contact_phone,
        }
        value = fields.get(reference.source_field)
        if reference.source_id != str(profile.id) or value is None:
            raise InvalidDraftForExport(
                "The CV draft references missing candidate profile information."
            )
    elif reference.source_type == "skill":
        item = session.get(CandidateSkill, UUID(reference.source_id))
        value = item.name if item and item.candidate_profile_id == profile.id else None
        if reference.source_field != "name":
            value = None
    elif reference.source_type == "work_experience":
        item = session.get(CandidateWorkExperience, UUID(reference.source_id))
        if item is None or item.candidate_profile_id != profile.id:
            value = None
        else:
            values = {
                "job_title": item.job_title,
                "employer": item.employer,
                "dates": _work_dates(item),
                "description": item.description,
            }
            if reference.source_field.startswith("achievements["):
                try:
                    index = int(
                        reference.source_field.removeprefix("achievements[").removesuffix("]")
                    )
                    value = item.achievements[index]
                except (ValueError, IndexError):
                    value = None
            else:
                value = values.get(reference.source_field)
    elif reference.source_type == "education":
        item = session.get(CandidateEducation, UUID(reference.source_id))
        if item is None or item.candidate_profile_id != profile.id:
            value = None
        else:
            values = {
                "qualification": item.qualification,
                "institution": item.institution,
                "field_of_study": item.field_of_study,
                "start_date": item.start_date.isoformat(),
                "end_date": item.end_date.isoformat() if item.end_date else None,
            }
            value = values.get(reference.source_field)
    else:
        item = session.get(CandidateEvidence, UUID(reference.source_id))
        if (
            item is None
            or item.candidate_profile_id != profile.id
            or item.confirmation_status != "candidate_confirmed"
            or reference.evidence_confirmation_status != "candidate_confirmed"
            or reference.source_document_id != str(item.source_document_id)
            or reference.source_excerpt != item.source_excerpt
            or reference.source_start != item.source_start
            or reference.source_end != item.source_end
        ):
            value = None
        else:
            value = item.description if reference.source_field == "description" else None

    if value is None or _clean(str(value)) != _clean(reference.claim_text):
        raise InvalidDraftForExport(
            "The CV draft contains a source reference that no longer matches "
            "candidate data or confirmed evidence."
        )


def validate_saved_draft(
    session: Session,
    profile_id: UUID,
    draft: TailoredCvDraft,
) -> tuple[CVStructuredContent, str]:
    if draft.candidate_profile_id != profile_id:
        raise InvalidDraftForExport("CV draft does not belong to this candidate profile.")
    if draft.status != "generated" or not draft.generated_content or not draft.structured_content:
        raise InvalidDraftForExport("Only successfully generated CV drafts can be exported.")
    validation_results = draft.validation_results
    if (
        not isinstance(validation_results, dict)
        or validation_results.get("is_valid") is not True
        or validation_results.get("errors")
    ):
        raise InvalidDraftForExport("This CV draft did not pass content validation.")

    try:
        content = CVStructuredContent.model_validate(draft.structured_content)
        references = [
            CVSourceReference.model_validate(item)
            for item in draft.source_references
        ]
    except (TypeError, ValueError) as exc:
        raise InvalidDraftForExport("The saved CV draft has invalid structured content.") from exc

    rendered = render_cv(content)
    errors = validate_cv_content(content, rendered)
    if errors or rendered != draft.generated_content:
        raise InvalidDraftForExport(
            "The saved CV draft no longer satisfies CV content validation."
        )

    profile = session.get(CandidateProfile, profile_id)
    if profile is None:
        raise InvalidDraftForExport("Candidate profile no longer exists.")
    for reference in references:
        try:
            _assert_reference_matches(reference, profile, session)
        except (ValueError, TypeError) as exc:
            if isinstance(exc, InvalidDraftForExport):
                raise
            raise InvalidDraftForExport(
                "The CV draft contains an invalid source reference."
            ) from exc
    return content, rendered


def _paragraph_text(value: str) -> str:
    return escape(value).replace("\n", "<br/>")


def _register_pdf_fonts() -> tuple[str, str]:
    font_dir = Path(__import__("reportlab").__file__).parent / "fonts"
    regular = font_dir / "Vera.ttf"
    bold = font_dir / "VeraBd.ttf"
    if not regular.is_file() or not bold.is_file():
        raise RuntimeError("Bundled Unicode-capable ReportLab fonts are unavailable.")
    if "JobPilotVera" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("JobPilotVera", str(regular)))
        pdfmetrics.registerFont(TTFont("JobPilotVera-Bold", str(bold)))
    return "JobPilotVera", "JobPilotVera-Bold"


class _PageCountCanvas(Canvas):
    def __init__(self, *args, page_count_holder: list[int], **kwargs) -> None:
        super().__init__(*args, **kwargs)
        self.page_count_holder = page_count_holder

    def save(self) -> None:
        self.page_count_holder.append(max(1, self.getPageNumber() - 1))
        super().save()


def render_pdf(content: CVStructuredContent) -> RenderedCv:
    regular_font, bold_font = _register_pdf_fonts()
    styles = getSampleStyleSheet()
    body = ParagraphStyle(
        "JobPilotBody",
        parent=styles["Normal"],
        fontName=regular_font,
        fontSize=9.5,
        leading=12.2,
        textColor=colors.black,
        alignment=TA_LEFT,
        spaceAfter=3,
    )
    heading = ParagraphStyle(
        "JobPilotHeading",
        parent=body,
        fontName=bold_font,
        fontSize=11,
        leading=14,
        spaceBefore=8,
        spaceAfter=4,
        keepWithNext=True,
    )
    name_style = ParagraphStyle(
        "JobPilotName",
        parent=heading,
        fontSize=19,
        leading=22,
        spaceBefore=0,
        spaceAfter=2,
    )
    subheading = ParagraphStyle(
        "JobPilotSubheading",
        parent=body,
        fontSize=10.5,
        leading=13,
        spaceAfter=6,
    )
    role_style = ParagraphStyle(
        "JobPilotRole",
        parent=body,
        fontName=bold_font,
        spaceBefore=4,
        spaceAfter=2,
        keepWithNext=True,
    )
    bullet_style = ParagraphStyle(
        "JobPilotBullet",
        parent=body,
        leftIndent=14,
        firstLineIndent=-9,
        spaceAfter=2,
    )

    story = [
        Paragraph(_paragraph_text(content.candidate_name), name_style),
        Paragraph(_paragraph_text(content.professional_heading), subheading),
    ]
    for section in content.sections:
        story.append(Paragraph(_paragraph_text(section.heading), heading))
        if section.heading == "Professional Experience":
            for role in section.roles:
                role_heading = Paragraph(
                    _paragraph_text(
                        f"{role.job_title} at {role.employer} | {role.dates}"
                    ),
                    role_style,
                )
                story.append(role_heading)
                story.extend(
                    Paragraph(
                        f"• {_paragraph_text(item.text)}",
                        bullet_style,
                    )
                    for item in role.bullets
                )
        else:
            for claim in section.claims:
                story.append(Paragraph(_paragraph_text(claim.text), body))
        story.append(Spacer(1, 2))

    output = BytesIO()
    page_count: list[int] = []

    def canvas_maker(*args, **kwargs) -> _PageCountCanvas:
        return _PageCountCanvas(*args, page_count_holder=page_count, **kwargs)

    document = SimpleDocTemplate(
        output,
        pagesize=letter,
        rightMargin=0.68 * inch,
        leftMargin=0.68 * inch,
        topMargin=0.58 * inch,
        bottomMargin=0.58 * inch,
        title="JobPilot AI CV",
        author="JobPilot AI",
    )
    document.build(story, canvasmaker=canvas_maker)
    return RenderedCv(
        content=output.getvalue(),
        page_count=page_count[0],
        media_type=MEDIA_TYPES["pdf"],
    )


def render_docx(content: CVStructuredContent) -> RenderedCv:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.58)
    section.bottom_margin = Inches(0.58)
    section.left_margin = Inches(0.68)
    section.right_margin = Inches(0.68)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(3)

    name = document.add_paragraph()
    name.paragraph_format.space_after = Pt(1)
    name.alignment = WD_ALIGN_PARAGRAPH.LEFT
    name_run = name.add_run(content.candidate_name)
    name_run.bold = True
    name_run.font.name = "Calibri"
    name_run.font.size = Pt(19)
    name_run.font.color.rgb = RGBColor(0, 0, 0)

    heading = document.add_paragraph()
    heading.paragraph_format.space_after = Pt(6)
    heading_run = heading.add_run(content.professional_heading)
    heading_run.font.name = "Calibri"
    heading_run.font.size = Pt(10.5)
    heading_run.font.color.rgb = RGBColor(0, 0, 0)

    for cv_section in content.sections:
        section_heading = document.add_paragraph()
        section_heading.paragraph_format.keep_with_next = True
        section_heading.paragraph_format.space_before = Pt(8)
        section_heading.paragraph_format.space_after = Pt(4)
        heading_run = section_heading.add_run(cv_section.heading)
        heading_run.bold = True
        heading_run.font.name = "Calibri"
        heading_run.font.size = Pt(11)
        heading_run.font.color.rgb = RGBColor(0, 0, 0)

        if cv_section.heading == "Professional Experience":
            for role in cv_section.roles:
                role_paragraph = document.add_paragraph()
                role_paragraph.paragraph_format.keep_with_next = True
                role_paragraph.paragraph_format.space_before = Pt(4)
                role_run = role_paragraph.add_run(
                    f"{role.job_title} at {role.employer} | {role.dates}"
                )
                role_run.bold = True
                role_run.font.name = "Calibri"
                role_run.font.color.rgb = RGBColor(0, 0, 0)
                for bullet in role.bullets:
                    paragraph = document.add_paragraph(style="List Bullet")
                    paragraph.paragraph_format.keep_together = True
                    paragraph.paragraph_format.space_after = Pt(2)
                    paragraph.add_run(bullet.text)
        else:
            for claim in cv_section.claims:
                paragraph = document.add_paragraph(claim.text)
                paragraph.paragraph_format.keep_together = True

    output = BytesIO()
    document.save(output)
    return RenderedCv(
        content=output.getvalue(),
        page_count=None,
        media_type=MEDIA_TYPES["docx"],
    )


def render_cv_export(
    content: CVStructuredContent,
    export_format: ExportFormat,
) -> RenderedCv:
    if export_format == "pdf":
        return render_pdf(content)
    if export_format == "docx":
        return render_docx(content)
    raise ValueError("Unsupported CV export format")


def create_cv_export(
    session: Session,
    storage: DocumentStorage,
    profile_id: UUID,
    draft: TailoredCvDraft,
    export_format: ExportFormat,
) -> TailoredCvExport:
    content, rendered_text = validate_saved_draft(session, profile_id, draft)
    rendered = render_cv_export(content, export_format)
    export_id = uuid4()
    extension = export_format
    filename = f"jobpilot-cv-{export_id.hex}.{extension}"
    try:
        stored = storage.store(BytesIO(rendered.content), suffix=f".{extension}")
    except (OSError, ValueError) as exc:
        raise ExportFileUnavailable("CV export could not be stored securely.") from exc

    document = CandidateDocument(
        candidate_profile_id=profile_id,
        document_kind="tailored_cv",
        provenance="system_generated",
        original_filename=filename,
        media_type=rendered.media_type,
        storage_key=stored.storage_key,
        size_bytes=stored.size_bytes,
        sha256=stored.sha256,
        target_job_title=None,
    )
    session.add(document)
    try:
        session.flush()
        export = TailoredCvExport(
            id=export_id,
            candidate_profile_id=profile_id,
            cv_draft_id=draft.id,
            document_id=document.id,
            export_format=export_format,
            page_count=rendered.page_count,
        )
        session.add(export)
        session.commit()
        session.refresh(export)
        return export
    except Exception:
        session.rollback()
        storage.remove(stored.storage_key)
        raise


def load_export_file(
    storage: DocumentStorage,
    document: CandidateDocument,
) -> Path:
    try:
        path = storage.resolve(document.storage_key)
    except DocumentStorageError as exc:
        raise ExportFileUnavailable("CV export file is unavailable.") from exc
    if not path.is_file():
        raise ExportFileUnavailable("CV export file is unavailable.")
    return path
